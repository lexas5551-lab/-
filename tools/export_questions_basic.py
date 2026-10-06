from openpyxl import load_workbook
from pathlib import Path
import json
SRC=Path('source/Исходный_банк_145_вопросов.xlsx')
wb=load_workbook(SRC,data_only=False)
items=[]
for domain,sheet,maxrow,opt_cols,key_col in [('ГНВП','Банк ГНВП 120',121,range(3,9),9),('ГОР','Банк ГОР 25',26,range(3,8),8)]:
    ws=wb[sheet]
    for r in range(2,maxrow+1):
        n=int(ws.cell(r,1).value); q=str(ws.cell(r,2).value)
        opts=[]
        for i,c in enumerate(opt_cols):
            v=ws.cell(r,c).value
            if v is not None and str(v).strip()!='': opts.append({'id':'ABCDEF'[i],'text':str(v)})
        key=[x.strip() for x in str(ws.cell(r,key_col).value).split(',')]
        items.append({'id':('GNVP' if domain=='ГНВП' else 'GOR')+f'-{n}','domain':domain,'number':n,'question':q,'options':opts,'correct':key})
print(json.dumps(items,ensure_ascii=False,indent=2))
